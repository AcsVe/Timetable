"""Reports: numbers, filters, Excel/PDF output, logo."""
import io
import uuid

import pytest
from openpyxl import load_workbook
from PIL import Image


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (60, 40), "#1f4e8c").save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture
def rep(school):
    """School with PE boys/girls groups, two lessons placed, a school name and a logo."""
    a, s = school["api"], school
    a.ok("post", "/api/school", {"name_ar": "مدرسة الاختبار", "name_en": "Test School"})
    d = a.ok("post", "/api/divisions", {"section_id": s["s5a"], "name": "أولاد/بنات"})
    boys = a.ok("post", "/api/groups", {"division_id": d["id"], "name_ar": "أولاد"})["id"]
    girls = a.ok("post", "/api/groups", {"division_id": d["id"], "name_ar": "بنات"})["id"]

    def lesson(subject, teacher, targets, ppw):
        return a.ok("post", "/api/lessons", {"timetable_id": s["tt"], "subject_id": s[subject], "periods_per_week": ppw,
                                             "teachers": [s[teacher]], "targets": targets})
    math = lesson("math", "t1", [{"section_id": s["s5a"]}], 4)
    lesson("math", "t1", [{"section_id": s["s10a"]}], 5)
    lesson("pe", "t_pe_m", [{"section_id": s["s5a"], "group_id": boys}], 2)
    lesson("pe", "t_pe_f", [{"section_id": s["s5a"], "group_id": girls}], 2)
    for i, c in enumerate(math["cards"][:2]):
        a.ok("patch", f"/api/cards/{c['id']}", {"version": c["version"], "weekday_id": s["days"][7], "period_no": i + 1})
    r = a.c.post("/api/school/logo", data={"file": (io.BytesIO(_png()), "logo.png")},
                 headers={"Idempotency-Key": str(uuid.uuid4())}, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    return s


def url(s, kind, **q):
    qs = "&".join(f"{k}={v}" for k, v in q.items())
    return f"/api/timetables/{s['tt']}/reports/{kind}?{qs}"


def test_section_coverage_counts_parallel_groups_once(rep):
    data = rep["api"].ok("get", url(rep, "stats-sections"))
    rows = {r[0]: r for r in data["tables"][0]["rows"]}
    row = rows["الخامس / أ"]
    # 4 math + PE: boys 2 ∥ girls 2 → 2, not 4
    assert row[3] == 6 and row[4] == 2 and row[2] == 30   # required, placed, 5 days × 6 periods


def test_teacher_sections_matrix_and_stage_filter(rep):
    t = rep["api"].ok("get", url(rep, "teacher-sections"))["tables"][0]
    assert t["columns"][0] == "المعلم" and "العاشر / أ" in t["columns"]
    ahmad = next(r for r in t["rows"] if r[0] == "أحمد")
    assert ahmad[-1] == 9
    basic_only = rep["api"].ok("get", url(rep, "teacher-sections", stage_id=rep["basic"]))["tables"][0]
    assert "العاشر / أ" not in basic_only["columns"]
    assert next(r for r in basic_only["rows"] if r[0] == "أحمد")[-1] == 4


def test_teacher_stats_against_target(rep):
    t = rep["api"].ok("get", url(rep, "stats-teachers", teacher_id=rep["t1"]))["tables"][0]
    (row,) = t["rows"]
    assert row[4:9] == [9, 2, 20, -11, 0.45]


def test_section_timetable_grid(rep):
    g = rep["api"].ok("get", url(rep, "section-timetable", section_id=rep["s5a"]))["grids"][0]
    assert g["title"] == "الشعبة: الخامس / أ"
    assert g["periods"][0] == {"no": 1, "time": "08:00–08:45"}
    cell = next(c for c in g["cells"] if c["day"] == 0 and c["period"] == 1)
    assert cell["entries"][0][0] == "الرياضيات"


def test_english_report(rep):
    data = rep["api"].ok("get", url(rep, "stats-subjects", lang="en"))
    assert data["title"] == "Subject statistics" and data["school_name"] == "Test School"


def test_xlsx_is_rtl_with_logo_and_real_numbers(rep):
    r = rep["api"].get(url(rep, "stats-teachers", format="xlsx"))
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb.active
    assert ws.sheet_view.rightToLeft is True
    assert len(ws._images) == 1
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert any(v == "مدرسة الاختبار" for row in values for v in row)
    header_row = next(i for i, row in enumerate(values) if "المعلم" in row)
    assert isinstance(values[header_row + 1][4], int)   # a number, not text


def test_xlsx_one_sheet_per_timetable(rep):
    r = rep["api"].get(url(rep, "section-timetable", format="xlsx", stage_id=rep["basic"]))
    wb = load_workbook(io.BytesIO(r.data))
    assert len(wb.sheetnames) == 2      # 5/A and 5/B


def test_pdf(rep):
    for kind in ("section-timetable", "teacher-timetable", "stats-sections", "teacher-subjects"):
        for lang in ("ar", "en"):
            r = rep["api"].get(url(rep, kind, format="pdf", lang=lang))
            assert r.status_code == 200 and r.data.startswith(b"%PDF") and len(r.data) > 5000, kind


def test_logo_rules(rep):
    a = rep["api"]
    assert a.get("/school-logo").status_code == 200
    r = a.c.post("/api/school/logo", data={"file": (io.BytesIO(b"GIF89a....."), "x.gif")},
                 headers={"Idempotency-Key": str(uuid.uuid4())}, content_type="multipart/form-data")
    assert r.status_code == 400
    school = a.ok("get", "/api/school")["items"][0]
    assert "logo_data" not in school and school["logo_path"].startswith("/school-logo")
    a.ok("delete", "/api/school/logo")
    assert a.get("/school-logo").status_code == 404


def test_unknown_kind_and_bad_filter(rep):
    assert rep["api"].get(url(rep, "nope")).status_code == 404
    assert rep["api"].get(url(rep, "stats-teachers", teacher_id="xyz")).status_code == 400


def test_class_teacher_subtitle_and_layouts(rep):
    a = rep["api"]
    sec = a.ok("get", f"/api/sections/{rep['s5a']}")
    a.ok("patch", f"/api/sections/{sec['id']}", {"version": sec["version"], "class_teacher_id": rep["t2"]})
    g = a.ok("get", url(rep, "section-timetable", section_id=rep["s5a"]))["grids"][0]
    assert g["subtitle"] == "مربي الصف: سارة"
    for layout, first_header in (("rows", "اليوم"), ("cols", "الحصة")):
        r = a.get(url(rep, "section-timetable", section_id=rep["s5a"], format="xlsx", layout=layout))
        ws = load_workbook(io.BytesIO(r.data)).active
        assert first_header in [c.value for row in ws.iter_rows() for c in row]
        r = a.get(url(rep, "section-timetable", section_id=rep["s5a"], format="pdf", layout=layout))
        assert r.data.startswith(b"%PDF")


def test_deleting_a_class_teacher_is_blocked(rep):
    a = rep["api"]
    sec = a.ok("get", f"/api/sections/{rep['s5b']}")
    a.ok("patch", f"/api/sections/{sec['id']}", {"version": sec["version"], "class_teacher_id": rep["t_pe_m"]})
    t = a.ok("get", f"/api/teachers/{rep['t_pe_m']}")
    r = a.delete(f"/api/teachers/{t['id']}?version={t['version']}")
    assert r.status_code == 409 and "section" in r.get_json()["details"]


def test_print_settings_letterhead_logo_size_and_preview(rep, monkeypatch):
    """Without the letterhead (or the logo) the page carries neither; larger sizes print larger text on one
    page per timetable; disposition=inline lets the browser show the PDF in the print preview."""
    import re
    from app.reports import pdf as P
    a = rep["api"]
    drawn = []
    real_rich = P.rich
    monkeypatch.setattr(P, "rich", lambda text, *x, **k: (drawn.append(str(text)), real_rich(text, *x, **k))[1])

    def pdf(**q):
        drawn.clear()
        r = a.get(url(rep, "section-timetable", format="pdf", section_id=rep["s5a"], **q))
        assert r.status_code == 200 and r.data[:4] == b"%PDF"
        return r

    images = lambda r: len(re.findall(rb"/Subtype\s*/Image", r.data))  # noqa: E731
    pages = lambda r: len(re.findall(rb"/Type\s*/Page[^s]", r.data))  # noqa: E731
    full = pdf()
    assert "مدرسة الاختبار" in drawn and images(full) == 1
    bare = pdf(header="0")
    assert "مدرسة الاختبار" not in drawn and images(bare) == 0
    nologo = pdf(logo="0")
    assert "مدرسة الاختبار" in drawn and images(nologo) == 0
    assert full.headers["Content-Disposition"].startswith("attachment")
    assert pdf(disposition="inline").headers["Content-Disposition"].startswith("inline")
    for size in ("large", "xlarge"):
        assert pages(pdf(size=size, header="0", lang="en")) == 1
    # Excel: no letterhead rows when asked
    from openpyxl import load_workbook as lw
    x = a.get(url(rep, "section-timetable", format="xlsx", section_id=rep["s5a"], header="0", size="large"))
    ws = lw(io.BytesIO(x.data)).worksheets[0]
    assert ws["A1"].value is None and not ws._images


def test_grid_fills_the_page_and_text_fits_its_columns():
    """The timetable takes the largest text size at which no word is cut and every row fits the page."""
    from app.reports.data import Grid
    from app.reports.pdf import _cell_height, _fit_size, _register, _width
    _register()
    heads = [f"{p}\n08:00–08:45" for p in range(1, 9)]
    cells = [[[["Computer Programming", "Yara Mardini"]] for _ in range(8)] for _ in range(5)]
    rows = [(d, c, [False] * 8) for d, c in zip(["الأحد", "الاثنين", "الثلاثاء", "الأربعاء", "الخميس"], cells)]
    small = _fit_size(heads, rows, 60, 60, 16, 80)
    wide = _fit_size(heads, rows, 140, 60, 16, 80)
    assert 6 <= small < wide <= 16
    assert _width("Programming", wide, True) <= 140 and _cell_height(cells[0][0], 140, wide) + 8 <= 80
    assert _cell_height([["Programming"]], 20, 10) is None            # a word wider than the column
    assert Grid(title="", days=[], periods=[]).cell_times == {}
