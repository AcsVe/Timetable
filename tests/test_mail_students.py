"""E-mail through Microsoft Graph (mocked HTTP), notify endpoints, students and their import."""
import io
import json
import uuid

import pytest
from openpyxl import Workbook, load_workbook

from app.notify import graph
from tests.helpers import make_user

TENANT = "11111111-2222-3333-4444-555555555555"
CLIENT = "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"


class FakeGraph:
    def __init__(self, fail_to=None, bad_secret=False):
        self.sent, self.tokens, self.fail_to, self.bad_secret = [], 0, fail_to, bad_secret

    def __call__(self, method, url, data=None, headers=None):
        if "login.microsoftonline.com" in url:
            self.tokens += 1
            assert TENANT in url and b"grant_type=client_credentials" in data
            if self.bad_secret:
                return 401, json.dumps({"error": "invalid_client", "error_description": "AADSTS7000215: Invalid client secret"}).encode()
            return 200, json.dumps({"access_token": "tok", "expires_in": 3600}).encode()
        assert url.startswith("https://graph.microsoft.com/v1.0/users/timetable%40school.test/sendMail")
        assert headers["Authorization"] == "Bearer tok"
        msg = json.loads(data)["message"]
        to = msg["toRecipients"][0]["emailAddress"]["address"]
        if to == self.fail_to:
            return 403, json.dumps({"error": {"code": "ErrorAccessDenied", "message": "Access is denied."}}).encode()
        self.sent.append({"to": to, "subject": msg["subject"], "html": msg["body"]["content"]})
        return 202, b""


@pytest.fixture
def fake(monkeypatch):
    f = FakeGraph()
    monkeypatch.setattr(graph, "_http", f)
    graph._token_cache.clear()
    return f


def configure(a):
    return a.ok("put", "/api/mail/settings", {"tenant_id": TENANT, "client_id": CLIENT, "client_secret": "s3cret",
                                              "sender": "timetable@school.test", "enabled": True})


def test_settings_hide_the_secret_and_validate(admin, fake):
    a = admin
    assert a.ok("get", "/api/mail/status") == {"configured": False}
    assert a.put("/api/mail/settings", json={"tenant_id": "x y"}).status_code == 400
    assert a.put("/api/mail/settings", json={"client_id": "nope"}).status_code == 400
    s = configure(a)
    assert s["configured"] and s["has_secret"] and "client_secret" not in s
    assert "s3cret" not in json.dumps(a.ok("get", "/api/settings"))
    assert a.put("/api/settings/mail:graph", json={"value": {}}).status_code == 400
    r = a.ok("post", "/api/mail/test", {"to": "me@school.test"})
    assert r["ok"] and fake.sent[0]["to"] == "me@school.test" and "dir='rtl'" in fake.sent[0]["html"]
    log = a.ok("get", "/api/mail/log")["items"]
    assert log[0]["status"] == "sent" and log[0]["kind"] == "test"


def test_bad_secret_is_reported_clearly(admin, monkeypatch):
    monkeypatch.setattr(graph, "_http", FakeGraph(bad_secret=True))
    graph._token_cache.clear()
    configure(admin)
    r = admin.ok("post", "/api/mail/test", {"to": "me@school.test"})
    assert not r["ok"] and "Invalid client secret" in r["error"]


def test_cover_notify_sends_mail_and_reports_failures(school, fake):
    a, s = school["api"], school
    configure(a)
    sun = s["days"][7]
    l = a.ok("post", "/api/lessons", {"timetable_id": s["tt"], "subject_id": s["math"], "periods_per_week": 2,
                                      "teachers": [s["t1"]], "targets": [{"section_id": s["s5a"]}]})
    for i, c in enumerate(l["cards"]):
        a.ok("patch", f"/api/cards/{c['id']}", {"version": c["version"], "weekday_id": sun, "period_no": i + 1})
    for tid, mail in ((s["t2"], "sara@school.test"), (s["t_pe_m"], "khaled@school.test")):
        t = a.ok("get", f"/api/teachers/{tid}")
        a.ok("patch", f"/api/teachers/{tid}", {"version": t["version"], "email": mail})
    a.ok("post", "/api/absences", {"teacher_id": s["t1"], "date_from": "2027-01-10", "date_to": "2027-01-10"})
    day = a.ok("get", f"/api/timetables/{s['tt']}/cover?date=2027-01-10")
    for n, sub in zip(day["needs"], (s["t2"], s["t_pe_m"])):
        a.ok("post", "/api/substitutions", {"timetable_id": s["tt"], "date": "2027-01-10", "card_id": n["card_id"],
                                            "period_no": n["period_no"], "original_teacher_id": s["t1"], "kind": "cover",
                                            "substitute_teacher_id": sub})
    fake.fail_to = "khaled@school.test"
    r = a.ok("post", f"/api/timetables/{s['tt']}/cover/notify", {"date": "2027-01-10"})
    assert (r["sent"], r["failed"]) == (1, 1) and fake.tokens == 1
    failed = next(x for x in r["results"] if x["email_status"] == "failed")
    assert "Mail.Send" in failed["error"]
    assert fake.sent[0]["subject"] == "حصص إشغال يوم الأحد 2027-01-10"
    notified = [n["substitution"]["notified_at"] for n in a.ok("get", f"/api/timetables/{s['tt']}/cover?date=2027-01-10")["needs"]]
    assert sum(1 for x in notified if x) == 1     # only the one who was reached


def test_invigilator_and_duty_notices(school, fake):
    a, s = school["api"], school
    configure(a)
    t = a.ok("get", f"/api/teachers/{s['t2']}")
    a.ok("patch", f"/api/teachers/{s['t2']}", {"version": t["version"], "email": "sara@school.test", "gender": "f"})
    a.ok("post", "/api/exam-sessions", {"term_id": s["term"], "exam_date": "2027-01-10", "subject_id": s["math"],
                                        "starts_at": "09:00", "ends_at": "10:00", "rooms": [{"teacher_ids": [s["t2"], s["t1"]]}]})
    r = a.ok("post", "/api/exam-sessions/notify", {"term_id": s["term"]})
    assert r["teachers"] == 2 and r["sent"] == 1 and r["no_email"] == 1
    assert "مراقبة واحدة" in fake.sent[0]["html"] and "الرياضيات" in fake.sent[0]["html"]
    a.ok("post", "/api/duty-assignments", {"term_id": s["term"], "duty_type": "الطابور الصباحي", "teacher_ids": [s["t2"]]})
    r = a.ok("post", "/api/duty-assignments/notify", {"term_id": s["term"], "email": False})
    assert r["teachers"] == 1 and r["results"][0]["email_status"] is None and len(fake.sent) == 1


# ---------------------------------------------------------------- students
def _xlsx(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(a, data, commit=True, **form):
    r = a.c.post(f"/api/import/{'commit' if commit else 'preview'}", data={"file": (io.BytesIO(data), "s.xlsx"), **form},
                 headers={"Idempotency-Key": str(uuid.uuid4())}, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    return r.get_json()


def test_import_students_creates_structure_groups_and_counts(school):
    a, s = school["api"], school
    d = a.ok("post", "/api/divisions", {"section_id": s["s5a"], "name": "بنين/بنات"})
    boys = a.ok("post", "/api/groups", {"division_id": d["id"], "name_ar": "بنين"})["id"]
    a.ok("post", "/api/groups", {"division_id": d["id"], "name_ar": "بنات"})
    data = _xlsx({"الطلبة": [
        ["رقم الطالب", "اسم الطالب", "الجنس", "المرحلة", "الصف", "الشعبة"],
        ["1", "محمد علي", "ذكر", "أساسي", "الخامس", "أ"],
        ["2", "ليان خالد", "أنثى", "أساسي", "الخامس", "أ"],
        ["3", "سامي حسن", "ذكور", "أساسي", "السادس", "ج"],          # new grade and section
        ["", "رنا يوسف", "اناث", "", "", "الخامس / ب"],
    ]})
    pre = _import(a, data, commit=False)
    assert pre["summary"]["students"]["created"] == 4 and pre["summary"]["grades"]["created"] == 1
    assert a.ok("get", "/api/students")["items"] == []           # preview saves nothing
    rep = _import(a, data)
    assert rep["summary"]["students"]["created"] == 4 and rep["error_count"] == 0
    st = {x["name_ar"]: x for x in a.ok("get", "/api/students")["items"]}
    assert st["محمد علي"]["group_ids"] == [boys]
    assert a.ok("get", f"/api/sections/{s['s5a']}")["student_count"] == 2
    # re-import: matched by number, moved to another section, nothing duplicated
    data2 = _xlsx({"Students": [["student no", "name", "grade", "section"], ["1", "محمد علي", "الخامس", "ب"]]})
    rep = _import(a, data2)
    assert rep["summary"]["students"] == {"created": 0, "updated": 1, "unchanged": 0}
    st = {x["name_ar"]: x for x in a.ok("get", "/api/students")["items"]}
    assert st["محمد علي"]["section_id"] == s["s5b"] and st["محمد علي"]["group_ids"] == []
    assert len(st) == 4


def test_teacher_sheet_sections_and_class_teacher(school):
    a, s = school["api"], school
    data = _xlsx({"المعلمون": [["اسم المعلم", "الشعب", "مربي الشعبة"], ["وفاء", "العاشر / أ، الخامس / ب", "العاشر / أ"]]})
    rep = _import(a, data)
    t = next(x for x in a.ok("get", "/api/teachers")["items"] if x["name_ar"] == "وفاء")
    assert set(t["stage_ids"]) == {s["basic"], s["secondary"]}
    assert a.ok("get", f"/api/sections/{s['s10a']}")["class_teacher_id"] == t["id"]
    assert rep["warning_count"] == 0
    # a chosen stage is given to teachers without a «المراحل» column
    data = _xlsx({"المعلمون": [["اسم المعلم"], ["نادر"]]})
    _import(a, data, stage_id=s["secondary"])
    t = next(x for x in a.ok("get", "/api/teachers")["items"] if x["name_ar"] == "نادر")
    assert t["stage_ids"] == [s["secondary"]]


def test_students_crud_move_lists_and_permissions(school, api):
    a, s = school["api"], school
    st = a.ok("post", "/api/students", {"section_id": s["s5a"], "name_ar": "أمل", "gender": "f", "student_no": "77"})
    assert a.post("/api/students", json={"section_id": s["s5a"], "name_ar": "x", "student_no": "77"}).status_code == 409
    d = a.ok("post", "/api/divisions", {"section_id": s["s5b"], "name": "بنين/بنات"})
    girls = a.ok("post", "/api/groups", {"division_id": d["id"], "name_ar": "بنات"})["id"]
    a.ok("post", "/api/students/move", {"section_id": s["s5b"], "items": [{"id": st["id"], "version": st["version"]}]})
    moved = a.ok("get", f"/api/students/{st['id']}")
    assert moved["section_id"] == s["s5b"] and moved["group_ids"] == [girls]
    rep = a.ok("get", f"/api/timetables/{s['tt']}/reports/student-lists")
    assert rep["tables"][0]["rows"][0][2] == "أمل" and "إناث: 1" in rep["tables"][0]["subtitle"]
    assert a.get(f"/api/timetables/{s['tt']}/reports/student-lists?format=pdf").data.startswith(b"%PDF")
    wb = load_workbook(io.BytesIO(a.get(f"/api/timetables/{s['tt']}/reports/student-lists?format=xlsx").data))
    assert len(wb.sheetnames) == 1
    # the section now has a student: it cannot be deleted
    sec = a.ok("get", f"/api/sections/{s['s5b']}")
    assert a.delete(f"/api/sections/{sec['id']}?version={sec['version']}").status_code == 409
    make_user("sec@school.test", role="stage_editor", stages=[s["secondary"]])
    a.logout()
    api.login("sec@school.test")
    assert api.post("/api/students", json={"section_id": s["s5a"], "name_ar": "y"}).status_code == 403
    assert api.post("/api/students", json={"section_id": s["s10a"], "name_ar": "y"}).status_code == 201
