"""Teacher absences, cover (حصص الإشغال), ranked substitutes, automatic assignment, notifications, reports."""
import io

import pytest
from openpyxl import load_workbook

from tests.helpers import make_user

SUNDAY = "2027-01-10"


def url(s, kind, **q):
    qs = "&".join(f"{k}={v}" for k, v in q.items())
    return f"/api/timetables/{s['tt']}/reports/{kind}?{qs}"


@pytest.fixture
def cov(school):
    """Sunday: Ahmad teaches 5A math in period 1; Sara 5A arabic in period 2 and 5B arabic in period 1;
    Khaled 5B math in period 3; Huda has no lesson on Sunday."""
    a, s = school["api"], school
    a.ok("post", "/api/school", {"name_ar": "مدرسة الاختبار"})
    sun = s["days"][7]

    def lesson(subject, teacher, section, period, ppw=3):
        l = a.ok("post", "/api/lessons", {"timetable_id": s["tt"], "subject_id": s[subject], "periods_per_week": ppw,
                                          "teachers": [s[t] for t in (teacher if isinstance(teacher, list) else [teacher])],
                                          "targets": [{"section_id": s[section]}]})
        c = l["cards"][0]
        a.ok("patch", f"/api/cards/{c['id']}", {"version": c["version"], "weekday_id": sun, "period_no": period})
        return l
    s["l_math"] = lesson("math", "t1", "s5a", 1)
    lesson("arabic", "t2", "s5a", 2)
    lesson("arabic", "t2", "s5b", 1)
    lesson("math", "t_pe_m", "s5b", 3)
    # Huda teaches on Monday only (so she is "not at school" on Sunday)
    l = a.ok("post", "/api/lessons", {"timetable_id": s["tt"], "subject_id": s["pe"], "periods_per_week": 1,
                                      "teachers": [s["t_pe_f"]], "targets": [{"section_id": s["s5b"]}]})
    c = l["cards"][0]
    a.ok("patch", f"/api/cards/{c['id']}", {"version": c["version"], "weekday_id": s["days"][1], "period_no": 1})
    return s


def day(s, d=SUNDAY):
    return s["api"].ok("get", f"/api/timetables/{s['tt']}/cover?date={d}")


def test_absence_creates_needs_and_ranked_candidates(cov):
    a, s = cov["api"], cov
    a.ok("post", "/api/absences", {"teacher_id": s["t1"], "date_from": SUNDAY, "date_to": SUNDAY, "reason": "إجازة مرضية"})
    d = day(s)
    assert d["weekday"]["name_ar"] == "الأحد" and d["summary"] == {"needs": 1, "decided": 0, "open": 1, "covered": 0,
                                                                     "absent_teachers": 1}
    need = d["needs"][0]
    assert need["period_no"] == 1 and need["subject"] == "الرياضيات" and need["starts_at"] == "08:00"
    cands = a.ok("get", f"/api/timetables/{s['tt']}/cover/candidates?date={SUNDAY}&key={need['key']}")["candidates"]
    names = [c["name"] for c in cands]
    assert "سارة" not in names            # teaching 5B in period 1
    assert "أحمد" not in names            # the absent teacher himself
    assert names[0] == "خالد"              # at school today, teaches maths
    huda = next(c for c in cands if c["name"] == "هدى")
    assert "ليس لديها حصص في هذا اليوم" in huda["warnings"]


def test_manual_cover_conflicts_and_edit(cov):
    a, s = cov["api"], cov
    a.ok("post", "/api/absences", {"teacher_id": s["t1"], "date_from": SUNDAY, "date_to": SUNDAY})
    need = day(s)["needs"][0]
    base = {"timetable_id": s["tt"], "date": SUNDAY, "card_id": need["card_id"], "period_no": 1,
            "original_teacher_id": s["t1"], "kind": "cover"}
    r = a.post("/api/substitutions", json={**base, "substitute_teacher_id": s["t2"]})
    assert r.status_code == 409 and "حصة في الوقت نفسه" in r.get_json()["details"]["message"]
    assert a.post("/api/substitutions", json={**base}).status_code == 400            # cover needs a teacher
    sub = a.ok("post", "/api/substitutions", {**base, "substitute_teacher_id": s["t_pe_m"]})
    # the same period cannot be decided twice
    assert a.post("/api/substitutions", json={**base, "substitute_teacher_id": s["t_pe_f"]}).status_code == 409
    sub = a.ok("patch", f"/api/substitutions/{sub['id']}", {"version": sub["version"], "kind": "cancel"})
    assert sub["substitute_teacher_id"] is None
    assert day(s)["summary"]["open"] == 0


def test_auto_assign_is_fair_and_handles_co_teachers(cov):
    a, s = cov["api"], cov
    # a co-taught lesson in period 4: Ahmad + Sara; Ahmad absent → no cover needed
    lab = a.ok("post", "/api/subjects", {"name_ar": "مختبر العلوم", "max_teachers_per_block": 2})["id"]
    l = a.ok("post", "/api/lessons", {"timetable_id": s["tt"], "subject_id": lab, "periods_per_week": 1,
                                      "teachers": [s["t1"], s["t2"]], "targets": [{"section_id": s["s10a"]}]})
    c = l["cards"][0]
    a.ok("patch", f"/api/cards/{c['id']}", {"version": c["version"], "weekday_id": s["days"][7], "period_no": 4})
    a.ok("post", "/api/absences", {"teacher_id": s["t1"], "date_from": SUNDAY, "date_to": SUNDAY})
    r = a.ok("post", f"/api/timetables/{s['tt']}/cover/auto", {"date": SUNDAY})
    assert r["created"] == 2 and r["unassigned"] == []
    d = day(s)
    kinds = sorted(n["substitution"]["kind"] for n in d["needs"])
    assert kinds == ["cover", "none"] and d["summary"]["open"] == 0
    assert all(n["substitution"]["auto"] for n in d["needs"])
    # clearing the automatic decisions only
    assert a.ok("post", f"/api/timetables/{s['tt']}/cover/clear", {"date": SUNDAY, "only_auto": True})["deleted"] == 2


def test_partial_absence_and_deleting_it_removes_decisions(cov):
    a, s = cov["api"], cov
    assert a.post("/api/absences", json={"teacher_id": s["t2"], "date_from": SUNDAY, "date_to": SUNDAY,
                                         "period_from": 2}).status_code == 400
    assert a.post("/api/absences", json={"teacher_id": s["t2"], "date_from": SUNDAY, "date_to": "2027-01-09"}).status_code == 400
    ab = a.ok("post", "/api/absences", {"teacher_id": s["t2"], "date_from": SUNDAY, "date_to": SUNDAY,
                                        "period_from": 2, "period_to": 3})
    d = day(s)
    assert [n["period_no"] for n in d["needs"]] == [2]          # her period-1 lesson is not affected
    a.ok("post", f"/api/timetables/{s['tt']}/cover/auto", {"date": SUNDAY})
    a.ok("delete", f"/api/absences/{ab['id']}?version={ab['version']}")
    assert day(s)["needs"] == []


def test_messages_notify_and_my_notifications(cov, api):
    a, s = cov["api"], cov
    uid = make_user("khaled@school.test", role="viewer")
    t = a.ok("get", f"/api/teachers/{s['t_pe_m']}")
    a.ok("patch", f"/api/teachers/{t['id']}", {"version": t["version"], "user_id": str(uid), "gender": "m"})
    a.ok("post", "/api/absences", {"teacher_id": s["t1"], "date_from": SUNDAY, "date_to": SUNDAY})
    need = day(s)["needs"][0]
    a.ok("post", "/api/substitutions", {"timetable_id": s["tt"], "date": SUNDAY, "card_id": need["card_id"], "period_no": 1,
                                        "original_teacher_id": s["t1"], "kind": "cover", "substitute_teacher_id": s["t_pe_m"]})
    r = a.ok("post", f"/api/timetables/{s['tt']}/cover/notify", {"date": SUNDAY})
    assert r["in_app"] == 1
    text = r["messages"][0]["text"].replace("\u200e", "").replace("\u200f", "")
    assert "لديك حصة واحدة للإشغال يوم الأحد 2027-01-10" in text and "بدلاً من المعلم أحمد" in text
    assert day(s)["needs"][0]["substitution"]["notified_at"]
    a.logout()
    api.login("khaled@school.test")
    mine = api.ok("get", "/api/notifications/mine")
    assert mine["unread"] == 1 and "الرياضيات" in mine["items"][0]["body"]["text"]
    api.ok("post", "/api/notifications/read", {})
    assert api.ok("get", "/api/notifications/mine")["unread"] == 0


def test_cover_reports(cov):
    a, s = cov["api"], cov
    a.ok("post", "/api/absences", {"teacher_id": s["t1"], "date_from": SUNDAY, "date_to": "2027-01-12", "reason": "دورة تدريبية"})
    a.ok("post", f"/api/timetables/{s['tt']}/cover/auto", {"date": SUNDAY})
    rep = a.ok("get", url(s, "cover-daily", date=SUNDAY))
    t1, t2 = rep["tables"]
    assert t1["rows"][0][4] == "أحمد" and t1["rows"][0][5] == "خالد" and t1["columns"][-1] == "توقيع البديل"
    assert t2["rows"][0][:3] == ["أحمد", "اليوم كاملاً", "دورة تدريبية"]
    stats = a.ok("get", url(s, "cover-stats", date_from="2027-01-01", date_to="2027-01-31"))["tables"][0]
    rows = {r[0]: r for r in stats["rows"]}
    assert rows["خالد"][1] == 1 and rows["أحمد"][2] == 3          # Sun–Tue = 3 school days absent
    log = a.ok("get", url(s, "absence-log", date_from="2027-01-01", date_to="2027-01-31"))["tables"][0]
    assert log["rows"][0][4] == 3
    free = a.ok("get", url(s, "free-teachers", date=SUNDAY))["grids"][0]
    assert free["days"] == ["الأحد"]
    p1 = next(c for c in free["cells"] if c["period"] == 1)["entries"][0][1]
    assert "أحمد" not in p1 and "خالد" not in p1                   # absent / covering in period 1
    for kind in ("cover-daily", "cover-stats", "absence-log"):
        assert a.get(url(s, kind, date=SUNDAY, format="pdf")).data.startswith(b"%PDF")
        wb = load_workbook(io.BytesIO(a.get(url(s, kind, date=SUNDAY, format="xlsx")).data))
        assert wb.sheetnames


def test_stage_editor_can_only_decide_for_own_stage(cov, api):
    a, s = cov["api"], cov
    a.ok("post", "/api/absences", {"teacher_id": s["t1"], "date_from": SUNDAY, "date_to": SUNDAY})
    make_user("sec@school.test", role="stage_editor", stages=[s["secondary"]])
    a.logout()
    api.login("sec@school.test")
    r = api.ok("post", f"/api/timetables/{s['tt']}/cover/auto", {"date": SUNDAY})
    assert r["created"] == 0            # Ahmad's 5A lesson belongs to the basic stage
