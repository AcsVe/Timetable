"""Excel (.xlsx) and CSV → Bundle.

A workbook may hold several sheets (teachers, subjects, rooms, sections, lessons);
each sheet is recognised by its name or, failing that, by its column headings, in
Arabic or English. A CSV file holds one table, recognised the same way (or by the
`kind` the user picked). The heading row may be preceded by a title row or two.
"""
from __future__ import annotations

import csv
import io

from app.errors import ApiError
from app.importing.bundle import Bundle
from app.importing.text import clean, gender, header_key, norm, split_list, to_int

KIND_LABEL = {"bells": ("التوقيت", "Bell times"), "teachers": ("المعلمون", "Teachers"), "subjects": ("المباحث", "Subjects"),
              "rooms": ("القاعات", "Rooms"), "sections": ("الشعب", "Sections"), "lessons": ("الدروس", "Lessons"),
              "students": ("الطلبة", "Students")}

SHEET_NAMES = {
    "teachers": ["المعلمون", "المعلمين", "معلمون", "المعلمات", "teachers", "staff"],
    "subjects": ["المباحث", "المواد", "مباحث", "subjects"],
    "rooms": ["القاعات", "الغرف", "قاعات", "rooms", "classrooms"],
    "sections": ["الشعب", "الصفوف والشعب", "شعب", "sections", "classes"],
    "lessons": ["الدروس", "الحصص", "التوزيع", "توزيع الحصص", "الإسناد", "lessons", "allocation", "cards"],
    "bells": ["التوقيت", "توقيت الحصص", "أوقات الحصص", "الجرس", "bells", "bell times", "periods"],
    "students": ["الطلبة", "الطلاب", "طلبة", "طلاب", "الطالبات", "قوائم الطلبة", "students", "pupils"],
}

FIELDS: dict[str, dict[str, list[str]]] = {
    "teachers": {
        "name": ["الاسم", "اسم المعلم", "المعلم", "المعلمة", "اسم المعلمة", "الاسم الكامل", "name", "teacher", "teacher name", "full name"],
        "name_en": ["الاسم بالإنجليزية", "الاسم بالانجليزي", "الاسم الإنجليزي", "name en", "english name"],
        "short": ["الاختصار", "الرمز", "المختصر", "short", "abbreviation", "code"],
        "email": ["البريد", "البريد الإلكتروني", "الإيميل", "email", "e mail", "mail"],
        "phone": ["الهاتف", "الجوال", "رقم الهاتف", "رقم الجوال", "phone", "mobile"],
        "title": ["اللقب", "title"],
        "gender": ["الجنس", "gender", "sex"],
        "target": ["النصاب", "النصاب الأسبوعي", "عدد الحصص المطلوبة", "target", "weekly load", "load"],
        "subjects": ["المباحث", "المواد", "subjects"],
        "stages": ["المراحل", "المرحلة", "stages", "stage"],
        "sections": ["الشعب", "الشعب التي يدرسها", "الصفوف والشعب", "الصفوف", "sections", "classes"],
        "class_of": ["مربي الشعبة", "مربي الصف", "مربية الصف", "مربي صف", "class teacher of", "homeroom of", "homeroom"],
    },
    "students": {
        "no": ["رقم الطالب", "الرقم المدرسي", "رقم الطالب في المدرسة", "الرقم", "الرقم الوطني", "رقم الهوية",
               "student no", "student number", "student id", "id"],
        "name": ["اسم الطالب", "اسم الطالبة", "الطالب", "الطالبة", "الاسم", "الاسم الكامل", "الاسم الرباعي", "name",
                 "student", "student name", "full name"],
        "name_en": ["الاسم بالإنجليزية", "الاسم بالانجليزي", "الاسم الإنجليزي", "name en", "english name"],
        "gender": ["الجنس", "النوع", "gender", "sex"],
        "stage": ["المرحلة", "stage"],
        "grade": ["الصف", "grade", "year", "class level"],
        "section": ["الشعبة", "الصف والشعبة", "الصف / الشعبة", "section", "class", "homeroom"],
        "groups": ["المجموعة", "المجموعات", "group", "groups"],
        "phone": ["هاتف ولي الأمر", "رقم ولي الأمر", "جوال ولي الأمر", "الهاتف", "الجوال", "guardian phone", "phone", "mobile"],
        "email": ["البريد", "البريد الإلكتروني", "email", "e mail"],
        "notes": ["ملاحظات", "notes"],
    },
    "subjects": {
        "name": ["الاسم", "المبحث", "اسم المبحث", "المادة", "اسم المادة", "name", "subject", "subject name"],
        "name_en": ["الاسم بالإنجليزية", "الاسم بالانجليزي", "name en", "english name"],
        "short": ["الاختصار", "الرمز", "short", "abbreviation", "code"],
        "short_en": ["الاختصار بالإنجليزية", "short en"],
        "color": ["اللون", "color", "colour"],
        "room_type": ["نوع القاعة المطلوبة", "نوع القاعة", "room type"],
    },
    "rooms": {
        "name": ["الاسم", "القاعة", "اسم القاعة", "الغرفة", "name", "room", "classroom"],
        "name_en": ["الاسم بالإنجليزية", "name en", "english name"],
        "short": ["الاختصار", "الرمز", "short", "code"],
        "room_type": ["النوع", "نوع القاعة", "type", "room type"],
        "capacity": ["السعة", "capacity", "seats"],
        "shared": ["مشتركة", "قاعة مشتركة", "shared"],
    },
    "sections": {
        "stage": ["المرحلة", "stage"],
        "grade": ["الصف", "grade", "year"],
        "name": ["الشعبة", "اسم الشعبة", "section", "class", "name"],
        "name_en": ["الاسم بالإنجليزية", "name en"],
        "student_count": ["عدد الطلبة", "عدد الطلاب", "الطلبة", "students", "student count"],
        "class_teacher": ["مربي الصف", "مربية الصف", "مربي الشعبة", "class teacher", "homeroom teacher"],
        "home_room": ["القاعة الصفية", "قاعة الشعبة", "home room", "homeroom"],
    },
    "lessons": {
        "subject": ["المبحث", "المادة", "subject"],
        "stage": ["المرحلة", "stage"],
        "grade": ["الصف", "grade", "year"],
        "sections": ["الشعبة", "الشعب", "section", "sections", "class", "classes"],
        "group": ["المجموعة", "group"],
        "teachers": ["المعلم", "المعلمون", "المعلمة", "teacher", "teachers"],
        "ppw": ["الحصص الأسبوعية", "عدد الحصص", "عدد الحصص الأسبوعية", "الحصص", "حصص", "periods per week", "periods", "ppw"],
        "duration": ["المدة", "طول البطاقة", "حصص متتالية", "duration", "periods per card", "length"],
        "room": ["القاعة المفضلة", "القاعة", "room", "preferred room"],
        "notes": ["ملاحظات", "notes"],
    },
    "bells": {
        "template": ["القالب", "اسم القالب", "قالب التوقيت", "template", "schedule"],
        "days": ["الأيام", "اليوم", "days", "day"],
        "grades": ["الصفوف", "الصف", "grades", "grade"],
        "period": ["الحصة", "رقم الحصة", "period"],
        "start": ["من", "البداية", "بداية", "start", "from"],
        "end": ["إلى", "الى", "النهاية", "نهاية", "end", "to"],
    },
}
_ALIASES = {kind: {header_key(a): f for f, al in fs.items() for a in al} for kind, fs in FIELDS.items()}
_SHEETS = {header_key(n): k for k, ns in SHEET_NAMES.items() for n in ns}
REQUIRED = {"bells": {"template", "period", "start", "end"}, "teachers": {"name"}, "subjects": {"name"}, "rooms": {"name"}, "sections": {"grade", "name"},
            "lessons": {"subject", "sections", "ppw"}, "students": {"name", "section"}}


def _find_header(rows: list[list], kind: str | None):
    """(row index, kind, {col: field}) of the first row that looks like a heading row."""
    kinds = [kind] if kind else ["bells", "lessons", "students", "sections", "teachers", "subjects", "rooms"]
    best = None
    for i, row in enumerate(rows[:10]):
        keys = [header_key(c) for c in row]
        for k in kinds:
            cols = {j: _ALIASES[k][key] for j, key in enumerate(keys) if key in _ALIASES[k]}
            fields = set(cols.values())
            if not REQUIRED[k] <= fields:
                continue
            score = len(fields) + (5 if k == "lessons" and "ppw" in fields else 0) + (2 if k == "students" else 0)
            if best is None or score > best[3]:
                best = (i, k, cols, score)
        if best:
            return best[:3]
    return None


def _rows_to_bundle(b: Bundle, rows: list[list], kind: str | None, where: str):
    found = _find_header(rows, kind)
    if not found:
        b.warn(f"لم تُعرف أعمدة الورقة «{where}»؛ تُجوهلت. استخدم عناوين القالب",
               f"Sheet '{where}': columns not recognised, skipped. Use the template headings", where)
        return
    hi, kind, cols = found
    for n, raw in enumerate(rows[hi + 1:], start=hi + 2):
        rec = {f: raw[j] if j < len(raw) else None for j, f in cols.items()}
        if not any(clean(v) for v in rec.values()):
            continue
        try:
            _ADD[kind](b, rec, where, n)
        except ValueError as e:
            b.error(f"قيمة غير صالحة «{e}»", f"Invalid value '{e}'", where, n)


def _num_field(rec, f, where, n, b, lo=0, hi=None):
    try:
        v = to_int(rec.get(f))
    except ValueError as e:
        raise ValueError(str(e))
    if v is not None and (v < lo or (hi is not None and v > hi)):
        raise ValueError(str(v))
    return v


def _opt(rec, f):
    return clean(rec.get(f)) or None


def _add_teacher(b, rec, where, n):
    name = clean(rec.get("name"))
    if not name:
        return b.error("اسم المعلم فارغ", "Teacher name is empty", where, n)
    b.teachers.append({"key": norm(name), "row": n, "where": where, "name": name, "name_en": _opt(rec, "name_en"),
                       "short": _opt(rec, "short"), "email": _opt(rec, "email"), "phone": _opt(rec, "phone"),
                       "title": _opt(rec, "title"), "gender": gender(rec.get("gender")),
                       "target": _num_field(rec, "target", where, n, b, 0, 60),
                       "subjects": split_list(rec.get("subjects")), "stages": split_list(rec.get("stages")),
                       "sections": split_list(rec.get("sections")), "class_of": _opt(rec, "class_of")})


def _add_student(b, rec, where, n):
    from app.importing.text import gender as _gender
    name, section = clean(rec.get("name")), clean(rec.get("section"))
    if not name or not section:
        return b.error("اسم الطالب والشعبة مطلوبان", "Student name and section are required", where, n)
    grade = _opt(rec, "grade")
    if not grade and "/" in section:
        grade, section = (clean(x) for x in section.rsplit("/", 1))
    if not grade:
        return b.error(f"حدِّد صف الطالب «{name}» (عمود «الصف» أو اكتب الشعبة بصيغة «السابع / أ»)",
                       f"Specify the grade of '{name}' (a 'Grade' column or 'Grade / Section')", where, n)
    try:
        g = _gender(rec.get("gender"))
    except ValueError:
        b.warn(f"قيمة الجنس «{clean(rec.get('gender'))}» غير مفهومة؛ تُركت فارغة", "Gender value not understood; left empty", where, n)
        g = None
    no = clean(rec.get("no")) or None
    b.students.append({"row": n, "where": where, "no": no[:40] if no else None, "name": name, "name_en": _opt(rec, "name_en"),
                       "gender": g, "stage": _opt(rec, "stage"), "grade": grade, "section": section,
                       "groups": split_list(rec.get("groups")), "phone": (_opt(rec, "phone") or "")[:30] or None,
                       "email": _opt(rec, "email"), "notes": _opt(rec, "notes")})


def _add_subject(b, rec, where, n):
    name = clean(rec.get("name"))
    if not name:
        return b.error("اسم المبحث فارغ", "Subject name is empty", where, n)
    color = _opt(rec, "color")
    if color and not color.startswith("#") and len(color) in (3, 6):
        color = "#" + color
    b.subjects.append({"key": norm(name), "row": n, "where": where, "name": name, "name_en": _opt(rec, "name_en"),
                       "short": _opt(rec, "short"), "short_en": _opt(rec, "short_en"), "color": color,
                       "room_type": _opt(rec, "room_type")})


def _add_room(b, rec, where, n):
    name = clean(rec.get("name"))
    if not name:
        return b.error("اسم القاعة فارغ", "Room name is empty", where, n)
    shared = norm(rec.get("shared"))
    b.rooms.append({"key": norm(name), "row": n, "where": where, "name": name, "name_en": _opt(rec, "name_en"),
                    "short": _opt(rec, "short"), "room_type": _opt(rec, "room_type"),
                    "capacity": _num_field(rec, "capacity", where, n, b, 0, 5000),
                    "shared": (shared in {"نعم", "1", "yes", "true", "x", "✓"}) if shared else None})


def _add_section(b, rec, where, n):
    grade, name = clean(rec.get("grade")), clean(rec.get("name"))
    if not grade or not name:
        return b.error("الصف والشعبة مطلوبان", "Grade and section are required", where, n)
    b.sections.append({"key": f"{norm(grade)}/{norm(name)}", "row": n, "where": where, "stage": _opt(rec, "stage"),
                       "grade": grade, "name": name, "name_en": _opt(rec, "name_en"),
                       "student_count": _num_field(rec, "student_count", where, n, b, 0, 1000),
                       "class_teacher": _opt(rec, "class_teacher"), "home_room": _opt(rec, "home_room")})


def _add_lesson(b, rec, where, n):
    subject = clean(rec.get("subject"))
    tokens = split_list(rec.get("sections"))
    ppw = _num_field(rec, "ppw", where, n, b, 1, 60)
    if not subject or not tokens or not ppw:
        return b.error("المبحث والشعبة وعدد الحصص الأسبوعية مطلوبة", "Subject, section and periods per week are required",
                       where, n)
    duration = _num_field(rec, "duration", where, n, b, 1, 4) or 1
    grade, stage, group = _opt(rec, "grade"), _opt(rec, "stage"), _opt(rec, "group")
    targets = []
    for tok in tokens:
        g, s = grade, tok
        if not g and "/" in tok:
            g, s = (clean(x) for x in tok.rsplit("/", 1))
        targets.append(({"stage": stage, "grade": g, "name": s, "text": tok if not grade else f"{grade} / {tok}"}, group))
    b.lessons.append({"key": None, "row": n, "where": where, "subject": subject, "targets": targets,
                      "teachers": split_list(rec.get("teachers")), "ppw": ppw, "duration": duration,
                      "room": _opt(rec, "room"), "notes": _opt(rec, "notes")})


BREAK_WORDS = {"استراحه", "الاستراحه", "فرصه", "الفرصه", "فسحه", "الفسحه", "break", "recess", "استراحة"}


def _add_bell(b, rec, where, n):
    from app.importing.text import time_text
    name = clean(rec.get("template"))
    if not name:
        return b.error("اسم القالب فارغ", "Template name is empty", where, n)
    p = clean(rec.get("period"))
    kind, no = ("break", None) if norm(p) in {norm(w) for w in BREAK_WORDS} else ("lesson", to_int(p))
    if kind == "lesson" and (no is None or not 1 <= no <= 20):
        raise ValueError(p)
    start, end = time_text(rec.get("start")), time_text(rec.get("end"))
    if not start or not end:
        return b.error("وقت البداية والنهاية مطلوبان", "Start and end times are required", where, n)
    tpl = next((x for x in b.bells if norm(x["template"]) == norm(name)), None)
    if tpl is None:
        tpl = {"template": name, "days": [], "grades": [], "slots": [], "where": where, "row": n}
        b.bells.append(tpl)
    for k in ("days", "grades"):
        for v in split_list(rec.get(k)):
            if v not in tpl[k]:
                tpl[k].append(v)
    tpl["slots"].append({"kind": kind, "no": no, "start": start, "end": end, "row": n, "where": where})


_ADD = {"bells": _add_bell, "teachers": _add_teacher, "subjects": _add_subject, "rooms": _add_room, "sections": _add_section,
        "lessons": _add_lesson, "students": _add_student}


def parse_xlsx(data: bytes, kind: str | None = None) -> Bundle:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ApiError("validation", 400, details={"reason": "bad_xlsx"},
                       message="تعذّر فتح ملف Excel؛ احفظه بصيغة ‎.xlsx وأعد المحاولة",
                       message_en="Could not open the Excel file; save it as .xlsx and try again")
    b = Bundle(source="xlsx")
    order = ["subjects", "rooms", "teachers", "sections", "students", "lessons", "bells"]
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or header_key(ws.title) in {header_key("تعليمات"), "instructions"}:
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        sheet_kind = _SHEETS.get(header_key(ws.title)) or (kind if len(wb.worksheets) == 1 else None)
        sheets.append((order.index(sheet_kind) if sheet_kind else 9, ws.title, rows, sheet_kind))
    for _, title, rows, sheet_kind in sorted(sheets, key=lambda x: x[0]):
        _rows_to_bundle(b, rows, sheet_kind, title)
    return b


def parse_csv(data: bytes, kind: str | None = None, name: str = "CSV") -> Bundle:
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    b = Bundle(source="csv")
    _rows_to_bundle(b, rows, kind, name)
    return b
