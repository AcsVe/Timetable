"""Excel export: real data cells (not pictures), right-to-left sheets for Arabic, logo in the header."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.reports.data import Grid, Report, Table, grid_matrix

THIN = Side(style="thin", color="B8C2D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E8C")
SUB_FILL = PatternFill("solid", fgColor="E8EEF7")
MISSING_FILL = PatternFill("solid", fgColor="EEEEEE")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_ROWS = 5
BLACK = Side(style="thin", color="000000")
THICK = Side(style="medium", color="000000")
PLAIN_BORDER = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
NO_FILL = PatternFill(fill_type=None)


class Look:
    """Cell styling for one report: 'plain' = black text and lines, no fills (black-and-white printing)."""

    def __init__(self, style: str):
        self.plain = style != "color"
        self.border = PLAIN_BORDER if self.plain else BORDER
        self.head_fill = NO_FILL if self.plain else HEAD_FILL
        self.head_font = Font(name="Arial", bold=True, color="000000" if self.plain else "FFFFFF")
        self.sub_fill = NO_FILL if self.plain else SUB_FILL
        self.missing_fill = NO_FILL if self.plain else MISSING_FILL
        self.title_color = "000000" if self.plain else "1F4E8C"
        self.muted = "000000" if self.plain else "555555"
        self.head_border = Border(left=BLACK, right=BLACK, top=THICK, bottom=THICK) if self.plain else BORDER


def _tint(hex_color: str, amount: float = 0.3) -> str:
    h = hex_color.lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda v: int(255 - (255 - v) * amount)  # noqa: E731
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name)[:28] or "Sheet"
    title, i = base, 2
    while title in used:
        title = f"{base[:25]} {i}"
        i += 1
    used.add(title)
    return title


def _header(ws, rep: Report, sub_title: str, width_cols: int, logo: bytes | None, look: "Look | None" = None):
    look = look or Look(rep.style)
    ws.sheet_view.rightToLeft = rep.rtl
    font = "Arial"
    first = 2 if logo else 1
    last = max(width_cols, first + 2)
    for row, text, size, bold in ((1, rep.school_name, 14, True), (2, rep.title, 13, True),
                                  (3, sub_title, 11, True),
                                  (4, " — ".join(x for x in (rep.timetable_name, rep.filters, rep.generated_at) if x), 9, False)):
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
        c = ws.cell(row=row, column=first, value=text)
        c.font = Font(name=font, size=size, bold=bold, color=look.title_color if row < 3 else ("000000" if look.plain else "333333"))
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    if logo:
        try:
            img = XLImage(io.BytesIO(logo))
            ratio = 70 / max(img.height, 1)
            img.height, img.width = 70, int(img.width * ratio)
            ws.add_image(img, "A1")
            ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, img.width / 7)
        except Exception:
            pass
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROWS + 1}:{HEADER_ROWS + 1}"


def _rich_cell(entries, muted="555555"):
    """Subject in bold, teacher/room under it in a smaller grey font (Excel rich text)."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    bold = InlineFont(rFont="Arial", b=True, sz=10)
    small = InlineFont(rFont="Arial", sz=8, color=muted)
    parts = []
    for e in entries:
        lines = [x for x in e if x]
        if not lines:
            continue
        if parts:
            parts.append(TextBlock(small, "\n"))
        parts.append(TextBlock(bold, lines[0]))
        for x in lines[1:]:
            parts.append(TextBlock(small, "\n" + x))
    return CellRichText(*parts) if parts else None


def _write_grid(ws, rep: Report, g: Grid, logo):
    look = Look(rep.style)
    corner, heads, rows = grid_matrix(g, rep.layout, rep.rtl, raw=True)
    _header(ws, rep, g.title + (f" — {g.subtitle}" if g.subtitle else ""), len(heads) + 1, logo, look)
    r0 = HEADER_ROWS + 1
    for j, text in enumerate([corner] + heads, start=1):
        c = ws.cell(row=r0, column=j, value=text)
        c.font, c.fill, c.alignment, c.border = look.head_font, look.head_fill, CENTER, look.head_border
    ws.row_dimensions[r0].height = 32 if rep.layout == "rows" else 18
    for i, (label, cells, missing) in enumerate(rows, start=1):
        row = r0 + i
        c = ws.cell(row=row, column=1, value=label)
        c.font, c.fill, c.alignment, c.border = Font(name="Arial", bold=True), look.sub_fill, CENTER, look.border
        max_lines = 2
        for j, (text, miss) in enumerate(zip(cells, missing), start=2):
            cell = ws.cell(row=row, column=j)
            cell.alignment, cell.border = CENTER, look.border
            if miss:
                if look.plain:
                    cell.value = "—"
                else:
                    cell.fill = look.missing_fill
                continue
            cell.font = Font(name="Arial", size=9)
            if text:
                cell.value = _rich_cell(text, look.muted)
                max_lines = max(max_lines, sum(len([x for x in e if x]) for e in text))
        ws.row_dimensions[row].height = 15 * max_lines
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 13)
    for j in range(len(heads)):
        ws.column_dimensions[get_column_letter(j + 2)].width = 16 if rep.layout == "rows" else 24
    end = r0 + len(rows) + 2
    if g.footer:
        ws.cell(row=end, column=1, value=g.footer).font = Font(name="Arial", italic=True)
        end += 1
    _write_signature(ws, rep, end + 1, len(heads) + 1)


def _write_signature(ws, rep: Report, row: int, width: int):
    if not rep.signature:
        return
    n = len(rep.signature)
    span = max(1, width // n)
    for i, label in enumerate(rep.signature):
        col = 1 + i * span
        ws.cell(row=row, column=col, value=label).font = Font(name="Arial", bold=True)
        ws.cell(row=row + 2, column=col, value="." * 24)


def _write_table(ws, rep: Report, t: Table, logo, start_row: int | None = None):
    look = Look(rep.style)
    if start_row is None:
        _header(ws, rep, " — ".join(x for x in (t.title, t.subtitle) if x), len(t.columns), logo, look)
        r0 = HEADER_ROWS + 1
    else:
        ws.cell(row=start_row, column=1, value=t.title or t.subtitle).font = Font(name="Arial", bold=True, size=12, color=look.title_color)
        r0 = start_row + 1
    if t.group_header:
        col = 1
        for label, span in t.group_header:
            c = ws.cell(row=r0, column=col, value=label or None)
            c.font, c.fill, c.alignment, c.border = look.head_font, look.head_fill, CENTER, look.head_border
            if span > 1:
                ws.merge_cells(start_row=r0, start_column=col, end_row=r0, end_column=col + span - 1)
            col += span
        r0 += 1
    for j, text in enumerate(t.columns, start=1):
        c = ws.cell(row=r0, column=j, value=text)
        c.font, c.fill, c.alignment, c.border = look.head_font, look.head_fill, CENTER, look.head_border
    rows = list(t.rows) + ([t.totals] if t.totals else [])
    for i, row in enumerate(rows, start=1):
        is_total = t.totals is not None and i == len(rows)
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            c.border = look.border
            c.font = Font(name="Arial", bold=is_total or j == 1, size=8 if t.compact else 11)
            if is_total:
                c.fill = look.sub_fill
                if look.plain:
                    c.border = Border(left=BLACK, right=BLACK, top=THICK, bottom=BLACK)
            color = t.cell_colors.get((i - 1, j - 1))
            if color and not look.plain and not is_total:
                c.fill = PatternFill("solid", fgColor=_tint(color))
            if (j - 1) in t.percent and isinstance(v, (int, float)):
                c.number_format = "0%"
            c.alignment = Alignment(horizontal="center" if ((j - 1) in t.numeric | t.percent or t.compact) else ("right" if rep.rtl else "left"),
                                    vertical="center", wrap_text=True)
        lines = max([1] + [str(v).count("\n") + 1 for v in row if isinstance(v, str)])
        if lines > 1:
            ws.row_dimensions[r0 + i].height = min(15 * lines, 400)
    for j, col in enumerate(t.columns, start=1):
        if t.compact and j > 1:
            ws.column_dimensions[get_column_letter(j)].width = 11
            continue
        longest = max([len(str(col))] + [max((len(x) for x in str(r[j - 1] or "").split("\n")), default=0) for r in rows])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, longest + 2), 60)
    if t.rows and not t.compact:
        ws.cell(row=r0 + len(rows) + 1, column=1, value=f"{'عدد الصفوف' if rep.rtl else 'Rows'}: {len(t.rows)}").font = Font(name="Arial", italic=True, size=9)
    if start_row is None:
        ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
        if t.rows and not t.group_header:
            ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(t.columns))}{r0 + len(t.rows)}"
    return r0 + len(rows) + 2


def render_xlsx(rep: Report, logo: bytes | None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for g in rep.grids:
        _write_grid(wb.create_sheet(_sheet_title(g.title.split(": ", 1)[-1], used)), rep, g, logo)
    if rep.tables and rep.kind == "student-lists":   # one sheet per section
        for t in rep.tables:
            _write_table(wb.create_sheet(_sheet_title(t.title.split(": ", 1)[-1], used)), rep, t, logo)
    elif rep.tables:
        ws = wb.create_sheet(_sheet_title(rep.title, used))
        end = _write_table(ws, rep, rep.tables[0], logo)
        for t in rep.tables[1:]:   # further tables of the same report go under the first one
            end = _write_table(ws, rep, t, logo, start_row=end + 2)
        for n in rep.notes:
            end += 2
            ws.cell(row=end, column=1, value=n).font = Font(name="Arial", italic=True)
        _write_signature(ws, rep, end + 3, len(rep.tables[0].columns))
    if not wb.sheetnames:
        ws = wb.create_sheet(_sheet_title(rep.title, used))
        _header(ws, rep, "لا توجد بيانات مطابقة" if rep.rtl else "No matching data", 4, logo)
    wb.properties.title = rep.title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
